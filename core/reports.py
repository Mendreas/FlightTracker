"""Exportações: Excel (openpyxl), PowerPoint (python-pptx) e relatórios HTML imprimíveis."""
from __future__ import annotations

import html
import time
from pathlib import Path

from .geo import fmt_date


# ------------------------------------------------------------------ Excel
def write_excel(path: str, sheets: list[tuple[str, list[str], list[list]]]):
    """sheets: [(nome, cabeçalhos, linhas)]."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(title=name[:31] or "Folha")
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F"); c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r) and r[i - 1] is not None]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(45, max(8, width))
        if rows:
            ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# ------------------------------------------------------------------ PPTX
def write_pptx_ias_app(path: str, hist: list[dict], caption: str, tol: float, airport: str = "LPPT"):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    prs = Presentation(); prs.slide_width = Inches(10); prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]

    def text(slide, s, x, y, w, h, size, bold=False, color="334155"):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = s; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        return tb
    s0 = prs.slides.add_slide(blank)
    text(s0, f"IAS na final (APP) — {airport}", 0.4, 0.22, 9.2, 0.5, 20, True, "0284C7")
    text(s0, caption, 0.4, 0.62, 9.2, 2.05, 9)
    text(s0, "Glossário", 0.4, 2.78, 9.2, 0.3, 12, True)
    gloss = ["n — Número de passagens com IAS válida ao CPA neste fix (após deduplicação pista/dia).",
             "Média — Soma das IAS dividida por n.",
             "Mediana — Valor que divide a distribuição ao meio (metade das observações abaixo, metade acima).",
             "Limite — Velocidade máxima publicada (AIP) para o fix.",
             f"Tolerância (±{tol:g} kt) — Igual ao campo da aba Velocidades APP; cada observação classifica-se face a limite ± tolerância.",
             "Cores nas barras — Verde: IAS entre limite−tol e limite+tol (conforme). Amarelo: abaixo de limite−tol. Vermelho: acima de limite+tol."]
    text(s0, "\n".join(gloss), 0.4, 2.98, 9.2, 2.5, 7.5)
    for h in hist:
        sl = prs.slides.add_slide(blank)
        text(sl, f"{h['point']} · RWY {h['rwy']}", 0.35, 0.2, 9.3, 0.4, 16, True, "0F172A")
        if not h["bins"]:
            text(sl, "Sem amostras com IAS ao CPA neste fix", 0.35, 0.6, 9.3, 0.4, 11); continue
        cd = CategoryChartData()
        cd.categories = [f"{b['lo']}" for b in h["bins"]]
        cd.add_series("Abaixo do limite", [b["y"] for b in h["bins"]])
        cd.add_series("Conforme", [b["g"] for b in h["bins"]])
        cd.add_series("Acima do limite", [b["r"] for b in h["bins"]])
        gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.35), Inches(0.55), Inches(9.3), Inches(4.0), cd)
        ch = gf.chart; ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        for ser, col in zip(ch.series, ["EAB308", "22C55E", "EF4444"]):
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
        ch.category_axis.tick_labels.font.size = Pt(8); ch.value_axis.tick_labels.font.size = Pt(8)
        text(sl, f"n={h['n']} · média {h['mean']:.0f} kt · mediana {h['median']:.0f} kt · limite {h['limit']} kt · tolerância ±{tol:g} kt",
             0.35, 4.6, 9.3, 0.4, 10)
    prs.save(path)


# ------------------------------------------------------------------ HTML (impressão)
CSS = """
body{font-family:Segoe UI,Arial,sans-serif;color:#0f172a;margin:24px;font-size:11pt}
h1{color:#1e3a8a;border-bottom:2px solid #1e3a8a;padding-bottom:6px;margin:0 0 4px}
h2{color:#1e3a8a;margin:22px 0 8px;font-size:14pt}h3{margin:16px 0 6px;font-size:12pt;color:#334155}
.sub{color:#64748b;font-size:9.5pt;margin-bottom:14px}
table{border-collapse:collapse;width:100%;font-size:9.5pt;margin:6px 0 12px}
th{background:#e2e8f0;text-align:left;padding:4px 6px;border:1px solid #cbd5e1}td{padding:3px 6px;border:1px solid #e2e8f0}
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:8px;margin:8px 0 14px}
.kpi{border:1px solid #cbd5e1;border-radius:6px;padding:6px 10px}.kpi b{display:block;font-size:15pt}.kpi span{font-size:8.5pt;color:#64748b;text-transform:uppercase}
.badge{display:inline-block;padding:2px 8px;border-radius:10px;font-size:9pt;font-weight:700}
.ok{background:#dcfce7;color:#166534}.bad{background:#fee2e2;color:#991b1b}
.note{font-size:9pt;color:#64748b;margin:4px 0 10px}.pb{page-break-before:always}
.insights{display:flex;flex-wrap:wrap;gap:6px;margin:6px 0 12px}.insights span{background:#eff6ff;border:1px solid #bfdbfe;border-radius:14px;padding:3px 10px;font-size:9pt}
footer{margin-top:26px;font-size:8.5pt;color:#94a3b8;border-top:1px solid #e2e8f0;padding-top:6px}
@media print{.noprint{display:none}}
"""


def esc(v) -> str:
    return html.escape("—" if v is None else str(v))


def fnum(v, d=0):
    if v is None or (isinstance(v, float) and v != v):
        return "—"
    s = f"{v:.{d}f}"
    return s.replace(".", ",")


def table(headers: list[str], rows: list[list], limit: int | None = None) -> str:
    rows = rows[:limit] if limit else rows
    h = "".join(f"<th>{esc(c)}</th>" for c in headers)
    b = "".join("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>" for r in rows)
    return f"<table><thead><tr>{h}</tr></thead><tbody>{b}</tbody></table>"


def kpi_grid(items: list[tuple[str, str]]) -> str:
    return '<div class="kpis">' + "".join(f'<div class="kpi"><b>{esc(v)}</b><span>{esc(k)}</span></div>' for k, v in items) + "</div>"


def svg_bars(values: list[float], labels: list[str], w=720, h=160, color="#2563eb", red: list[bool] | None = None) -> str:
    if not values:
        return '<p class="note">Sem dados.</p>'
    mx = max(values) or 1
    n = len(values); bw = (w - 40) / n
    parts = [f'<svg viewBox="0 0 {w} {h}" width="{w}" height="{h}" xmlns="http://www.w3.org/2000/svg" style="font-family:Arial;font-size:9px">']
    for i, (v, lb) in enumerate(zip(values, labels)):
        bh = (h - 40) * v / mx
        x = 20 + i * bw
        parts.append(f'<rect x="{x + 2:.1f}" y="{h - 24 - bh:.1f}" width="{bw - 4:.1f}" height="{bh:.1f}" fill="{"#dc2626" if red and red[i] else color}"/>')
        parts.append(f'<text x="{x + bw / 2:.1f}" y="{h - 28 - bh:.1f}" text-anchor="middle">{v:g}</text>')
        parts.append(f'<text x="{x + bw / 2:.1f}" y="{h - 10}" text-anchor="middle" fill="#64748b">{esc(lb)}</text>')
    parts.append("</svg>")
    return "".join(parts)


def svg_grouped(series: list[tuple[str, list[float], str]], labels: list[str], w=720, h=180) -> str:
    if not labels:
        return '<p class="note">Sem dados.</p>'
    mx = max((max(v) for _, v, _ in series if v), default=1) or 1
    n = len(labels); gw = (w - 40) / n; bw = (gw - 6) / max(1, len(series))
    p = [f'<svg viewBox="0 0 {w} {h}" width="{w}" height="{h}" xmlns="http://www.w3.org/2000/svg" style="font-family:Arial;font-size:9px">']
    for i, lb in enumerate(labels):
        for j, (name, vals, col) in enumerate(series):
            v = vals[i] if i < len(vals) else 0; bh = (h - 44) * v / mx
            x = 20 + i * gw + 3 + j * bw
            p.append(f'<rect x="{x:.1f}" y="{h - 28 - bh:.1f}" width="{bw - 1:.1f}" height="{bh:.1f}" fill="{col}"/>')
        p.append(f'<text x="{20 + i * gw + gw / 2:.1f}" y="{h - 14}" text-anchor="middle" fill="#64748b">{esc(lb)}</text>')
    lx = 20
    for name, _, col in series:
        p.append(f'<rect x="{lx}" y="{h - 8}" width="10" height="8" fill="{col}"/><text x="{lx + 14}" y="{h - 1}">{esc(name)}</text>'); lx += 14 + 7 * len(name) + 10
    p.append("</svg>")
    return "".join(p)


def wrap(title: str, subtitle: str, body: str, landscape=False) -> str:
    size = "@page{size:A4 landscape;margin:12mm}" if landscape else "@page{size:A4;margin:14mm}"
    return (f"<!DOCTYPE html><html lang=\"pt\"><head><meta charset=\"utf-8\"><title>{esc(title)}</title><style>{CSS}{size}</style></head>"
            f"<body><button class=\"noprint\" onclick=\"window.print()\" style=\"float:right;padding:6px 14px\">🖨 Imprimir / PDF</button>"
            f"<h1>{esc(title)}</h1><div class=\"sub\">{subtitle}</div>{body}"
            f"<footer>FlightTracker — NAV Portugal. Relatório automático gerado em {time.strftime('%d/%m/%Y %H:%M')}. Dados baseados em track radar.</footer></body></html>")


# ---------------- relatórios concretos
def report_flight(flight: dict, pts: list[dict], app_rec: dict | None, star_viol: list[dict], tol: float) -> str:
    alts = [p["alt"] for p in pts if p["alt"] is not None]; ias = [p["ias"] for p in pts if p["ias"]]
    stats = [("Data", flight["date_fmt"]), ("Pontos", str(len(pts))), ("Alt máx", f"{max(alts)} ft" if alts else "—"), ("Alt mín", f"{min(alts)} ft" if alts else "—"),
             ("IAS máx", f"{max(ias):.0f} kt" if ias else "—"), ("IAS méd", f"{sum(ias) / len(ias):.0f} kt" if ias else "—"),
             ("Início", pts[0]["tod"][:8] if pts else "—"), ("Fim", pts[-1]["tod"][:8] if pts else "—")]
    app_pts = [p for p in (app_rec or {}).get("points", []) if p["ias"] is not None]
    app_viol = [p for p in app_pts if not p["point"].upper().startswith("THR") and abs(p["deviation"]) > tol]
    badge = '<span class="badge ok">✔ Sem violações</span>' if not app_viol and not star_viol else \
        f'<span class="badge bad">⚠ {len(app_viol)} viol. APP · {len(star_viol)} viol. STAR</span>'
    paras = [f"{flight['csn']} — análise de velocidades para {flight['date_fmt']}."]
    if ias: paras.append(f"Velocidade máxima registada: {max(ias):.0f} kt. Média: {sum(ias) / len(ias):.0f} kt.")
    if app_pts:
        paras.append(f"Na fase de aproximação, foram avaliados {len(app_pts)} ponto(s): " + ", ".join(p["point"] for p in app_pts) + ".")
        above = [p for p in app_viol if p["deviation"] > 0]; below = [p for p in app_viol if p["deviation"] < 0]
        if above: paras.append(f"⚠ {len(above)} ponto(s) acima do limite: " + "; ".join(f"{p['point']} ({p['ias']:.0f}kt, lim. {p['limit']}kt, +{p['deviation']:.0f}kt)" for p in above))
        if below: paras.append(f"▼ {len(below)} ponto(s) abaixo do limite: " + "; ".join(f"{p['point']} ({p['ias']:.0f}kt, lim. {p['limit']}kt, {p['deviation']:.0f}kt)" for p in below))
        if app_viol:
            m = max(app_viol, key=lambda p: abs(p["deviation"]))
            paras.append(f"Desvio máximo: {m['point']} com {m['deviation']:+.0f}kt (lim. {m['limit']}kt).")
        else:
            paras.append(f"✔ Todos os pontos de aproximação dentro do limite (tolerância ±{tol:g}kt).")
    else:
        paras.append("Sem dados de velocidade de aproximação disponíveis.")
    if star_viol:
        wps = sorted({v["wp"] for v in star_viol})
        paras.append(f"Na fase STAR: {len(star_viol)} violação(es) em {len(wps)} waypoint(s) — {', '.join(wps)}. Excesso máximo: +{max(v['excess'] for v in star_viol):.0f}kt.")
    else:
        paras.append("✔ Sem violações de velocidade na fase STAR.")
    n = len(app_viol) + len(star_viol)
    paras.append("Conclusão: Voo sem violações de velocidade detectadas." if n == 0 else
                 f"Conclusão: Registadas {n} violação(es). Recomenda-se revisão do perfil de desaceleração.")
    body = badge + kpi_grid(stats) + "".join(f"<p>{esc(p)}</p>" for p in paras)
    # perfil IAS/altitude (SVG simples)
    prof = [(p["dist_thr"], p["ias"], p["alt"]) for p in pts if p.get("dist_thr") is not None and p["dist_thr"] <= 20]
    if prof:
        w, h = 720, 220; mxd = max(d for d, _, _ in prof) or 1; mxa = max((a for _, _, a in prof if a is not None), default=1) or 1
        mxi = max((i for _, i, _ in prof if i), default=1) or 1
        pi = " ".join(f"{20 + (w - 40) * (1 - d / mxd):.1f},{h - 30 - (h - 60) * (i / mxi):.1f}" for d, i, _ in prof if i)
        pa = " ".join(f"{20 + (w - 40) * (1 - d / mxd):.1f},{h - 30 - (h - 60) * (a / mxa):.1f}" for d, _, a in prof if a is not None)
        body += f"<h2>Perfil IAS / Altitude</h2><svg viewBox='0 0 {w} {h}' width='{w}' height='{h}' style='font-family:Arial;font-size:9px'>" \
                f"<polyline points='{pi}' fill='none' stroke='#0891b2' stroke-width='2'/><polyline points='{pa}' fill='none' stroke='#16a34a' stroke-width='1.5' stroke-dasharray='4 3'/>" \
                f"<text x='20' y='12' fill='#0891b2'>IAS (kt) — máx {mxi:.0f}</text><text x='200' y='12' fill='#16a34a'>Altitude (ft) — máx {mxa}</text>" \
                f"<text x='20' y='{h - 8}' fill='#64748b'>{mxd:.0f} NM</text><text x='{w - 40}' y='{h - 8}' fill='#64748b'>THR</text></svg>"
    if app_pts:
        body += "<h2>Velocidades APP</h2>" + table(["Pista", "Waypoint", "IAS", "Limite", "Desvio", "Avaliação"],
            [[f"RWY {app_rec['rwy']}", p["point"], f"{p['ias']:.0f}", p["limit"], f"{p['deviation']:+.0f}",
              "⚠ Acima" if p in app_viol and p["deviation"] > 0 else "▼ Abaixo" if p in app_viol else "✔ OK"] for p in app_pts])
    if star_viol:
        body += "<h2>Velocidades STAR</h2>" + table(["Hora", "Data", "Waypoint", "IAS", "Limite", "Excesso"],
            [[v["tod"], v["date"], v["wp"], f"{v['ias']:.0f}", v["limit"], f"+{v['excess']:.0f}"] for v in star_viol])
    if not app_pts and not star_viol:
        body += '<p class="note">Sem dados de violações registados para este voo.</p>'
    body += table(["TOD", "Alt ft", "IAS", "GS", "ROC", "Dist THR"], [[p["tod"], p["alt"], fnum(p["ias"]), fnum(p["gs"]), fnum(p["roc"]),
                                                                       fnum(p.get("dist_thr"), 1) + " NM"] for p in pts], 400) if pts else ""
    return wrap(f"Relatório {flight['csn']}", f"Relatório de Voo · FlightTracker · NAV Portugal · {flight['date_fmt']}", body)


def report_kpis(k: dict, airport: str) -> str:
    body = "<h2>Resumo Operacional</h2>" + kpi_grid([("Movimentos", k["total"]), ("Chegadas (ARR)", k["arr"]), ("Partidas (DEP)", k["dep"]), ("Sobrevoos (LVL)", k["lvl"]),
                                                     ("Dias de dados", k["days"]), ("Hora de pico", f"{k['peakHour']:02d}:00 ({k['peakN']} mov.)"), ("Período (UTC)", k["period"]), ("Amostras radar", k["pts"])])
    if k.get("approach"):
        a = k["approach"]
        body += "<h2>Análise de Aproximações</h2>" + kpi_grid([("Aterragens", a["n"]), ("Taxa", f"{a['rate']:.1f} aterr./h" if a["rate"] else "—"),
                                                                ("Estabilizadas", f"{a['stab']}/{a['n']}"), ("IAS@FAF méd.", fnum(a["iasFAF"]) + " kt"),
                                                                ("Tempo FAF→THR méd.", fnum(a["durS"]) + " s"), ("Sep. temporal méd.", fnum(a["sepMean"]) + " s"),
                                                                ("Sep. mín/máx", f"{fnum(a['sepMin'])} / {fnum(a['sepMax'])} s"),
                                                                ("Por pista", " · ".join(f"RWY {r}: {n}" for r, n in sorted(a["byRwy"].items())))])
    if k.get("landingRate"):
        lr = k["landingRate"]
        body += "<h2>Taxa de aterragem (rolling 60 min)</h2>" + kpi_grid([("Pico 60 min", f"{lr['peak']} aterr./h"), ("% capacidade (38)", f"{lr['pct']:.0f}%"), ("Média período", f"{lr['avg']} aterr./h")])
        body += svg_bars(lr["series"], lr["labels"], color="#f87171")
    if k.get("sepHist"):
        s = k["sepHist"]
        body += "<h2>Separação temporal entre aterragens consecutivas</h2>" + kpi_grid([("Média", fnum(s["mean"]) + " s"), ("Mediana", f"{s['median']:.0f} s"), ("Mínimo", f"{s['min']:.0f} s"), ("Máximo", f"{s['max']:.0f} s"), ("Bunching (<90s)", s["bunching"])])
        body += svg_bars(s["counts"], s["labels"], color="#f87171")
    if k.get("cdo"):
        c = k["cdo"]
        body += "<h2>CDO — Continuous Descent Operations (&lt; FL100)</h2>" + kpi_grid([("Taxa CDO", f"{c['pct']:.0f}%"), ("Aproximações", f"{c['count']}/{c['n']}"), ("Level-offs méd.", f"{c['levelOffs']:.1f}"), ("Tempo nível méd.", f"{c['levelSecs']:.0f} s")])
    if k.get("star"):
        s = k["star"]
        body += "<h2>Cumprimento de IAS em waypoints STAR (ARR)</h2>" + kpi_grid([("Passagens", s["passes"]), ("Conformes", s["comp"]), ("Violações", s["viol"]), ("Taxa de cumprimento", fnum(s["pct"], 1) + " %")])
        if s["top5"]:
            body += "<h3>Top 5 waypoints com mais violações</h3>" + table(["Waypoint", "Limite", "Violações", "Passagens", "%"], [[w["wp"], w["limit"], w["violations"], w["total"], f"{w['violations'] / w['total'] * 100:.1f}"] for w in s["top5"]])
        if s["operators"]:
            body += "<h3>Cumprimento IAS por operador (top 10 pior)</h3>" + table(["Operador", "Voos ARR", "Violações", "Inc./voo"], [[o["op"], o["arrFlights"], o["violations"], f"{o['rate'] * 100:.1f}%"] for o in s["operators"]])
    body += "<h2>Movimentos por hora (UTC)</h2>" + svg_grouped([("ARR", k["byArr"], "#f87171"), ("DEP", k["byDep"], "#22d3a0")], [f"{h:02d}" for h in range(24)])
    if len(k.get("byDay", [])) > 1:
        body += "<h2>Movimentos por dia</h2>" + table(["Dia", "ARR", "DEP", "LVL", "Total"], [[d["label"], d["arr"], d["dep"], d["lvl"], d["tot"]] for d in k["byDay"]])
    body += "<h2>Top 10 operadores (prefixo callsign)</h2>" + table(["Operador", "Movimentos"], [[o, n] for o, n in k["topOps"]])
    return wrap(f"Relatório Operacional — {airport}", f"Gerado em {time.strftime('%d/%m/%Y %H:%M')} · Dados: {k['callsigns']} callsigns", body)


def report_data_analysis(da: dict, airport: str) -> str:
    k = da["kpis"]; n = da["narrative"]; ch = da["charts"]; limit = da["filters"]["Limite crítico"]
    body = "<h2>Filtros</h2>" + table(["Filtro", "Valor"], [[a, b] for a, b in da["filters"].items()])
    body += "<h2>Síntese executiva</h2>" + "".join(f"<p>{esc(p)}</p>" for p in n["paragraphs"])
    body += "<h2>Destaques</h2><div class='insights'>" + "".join(f"<span>{esc(i)}</span>" for i in n["insights"]) + "</div>"
    body += "<h2>Indicadores principais</h2>" + kpi_grid([("Pares analisados", k["pairs"]), ("Pares abaixo do limite", k["crit"]), ("Críticos com DEP", k["critDep"]), ("Críticos sem DEP", k["critNoDep"]),
                                                          ("Menor separação", fnum(k["min"], 1) + " s"), ("Média", fnum(k["avg"], 1) + " s"), ("Mediana", fnum(k["median"], 1) + " s"), ("Hora mais crítica", k["critHour"]),
                                                          ("Hora mais movimentada", k["busyHour"]), ("Pista mais usada", k["mainRwy"]), ("DEP entre chegadas", k["depTotal"]), ("Movimentos ARR+DEP", k["movs"])])
    body += f"<h2>Distribuição de separações</h2><p class='note'>Histograma de pares por intervalo de separação (segundos). Barras a vermelho intersectam o limite de {limit} — concentração à esquerda indica sequências mais apertadas.</p>"
    body += svg_bars(ch["hist"]["total"], ch["hist"]["labels"], red=ch["hist"]["red"])
    hours = sorted(ch["byHour"])
    body += f"<h2>Pares por hora</h2><p class='note'>Total de pares precedente/seguidor vs críticos por hora UTC. A hora {k['critHour']} concentra o maior número de críticos.</p>"
    body += svg_grouped([("Total", [ch["byHour"][h]["total"] for h in hours], "#2563eb"), ("Críticos", [ch["byHour"][h]["crit"] for h in hours], "#dc2626")], [f"{h}h" for h in hours])
    body += "<h2>Comparação por pista</h2><p class='note'>Volume, críticos e DEP intercalado por RWY 02 e RWY 20 — permite comparar pressão operacional entre cabeceiras.</p>"
    body += table(["Pista", "Pares", "Críticos", "Com DEP"], [[f"RWY {r}", v["total"], v["crit"], v["dep"]] for r, v in ch["byRwy"].items()])
    body += "<h2>Sequências ARR/DEP</h2><p class='note'>Tipos de sequência (AA, ADA, ADDA, AD…A): AA sem DEP intercalado; ADA/ADDA com uma ou mais descolagens entre chegadas.</p>"
    body += svg_grouped([("Total", [ch["bySeq"][s]["total"] for s in ch["bySeq"]], "#2563eb"), ("Críticos", [ch["bySeq"][s]["crit"] for s in ch["bySeq"]], "#dc2626")], list(ch["bySeq"]))
    mh = ch["movHour"]
    body += "<h2>Movimentos por hora</h2><p class='note'>Perfil horário de movimentos ARR+DEP nos filtros activos — complementa a análise de pares com a carga operacional global.</p>"
    body += svg_bars([mh[h]["total"] for h in mh], [f"{h}h" for h in mh])
    body += "<h2 class='pb'>Companhias aéreas</h2>"
    if ch["companies"]:
        body += f"<p class='note'>Top 15 operadores por movimentos no período filtrado; {ch['companies'][0]['company']} lidera com {ch['companies'][0]['movements']} movimentos.</p>"
        body += svg_bars([c["movements"] for c in ch["companies"]], [c["company"] for c in ch["companies"]])
    if ch["topCritCompanies"]:
        body += "<h3>Taxa de ocorrências críticas</h3><p class='note'>Companhias com maior taxa de críticos por movimento — útil para identificar operadores com sequências mais apertadas.</p>"
        body += svg_bars([round(c["rate"] * 100, 1) for c in ch["topCritCompanies"]], [c["company"] for c in ch["topCritCompanies"]], color="#dc2626")
    body += "<h2 class='pb'>Factores operacionais</h2><h3>Posição do seguidor no DEP</h3><p class='note'>Distância do seguidor à THR no momento da descolagem intercalada. Valores baixos (&lt;3–4 NM) indicam DEP mais apertado e maior risco de sequência crítica.</p>"
    body += svg_bars(ch["depDist"]["total"], [f"{l} NM" for l in ch["depDist"]["labels"]], red=[True, True, False, False, False])
    body += "<h3>IAS do seguidor</h3><p class='note'>Distribuição de pares por classe de IAS do seguidor na THR — velocidades mais altas podem reduzir margem de recuperação em sequências apertadas.</p>"
    body += svg_grouped([("Total", ch["ias"]["total"], "#2563eb"), ("Críticos", ch["ias"]["crit"], "#dc2626")], ch["ias"]["labels"])
    t = da["tables"]
    body += "<h2 class='pb'>Anexos tabulares</h2><h3>Top 10 pares mais apertados</h3>" + table(["Pista", "Data", "Hora", "Precedente", "Seguidor", "Seq.", "Sep. s", "Sep. NM", "DEP"],
        [[f"RWY {r['rwy']}", r["date"], r["leadTod"], r["lead"], r["trail"], r["seq"], round(r["sep"]), r["sepNm"], r["firstDep"] or "Sem DEP"] for r in t["top10"]])
    body += "<h3>Resumo por hora</h3>" + table(["Hora", "Chegadas", "Pares", "Críticos", "DEP", "Menor sep.", "Média sep."],
        [[f"{h['hour']}h", h["arrivals"], h["pairs"], h["crit"], h["dep"], f"{h['minSep']} s", fnum(h["avgSep"], 1) + " s"] for h in t["hour"]], 24)
    body += "<h3>Companhias</h3>" + table(["Companhia", "Movimentos", "ARR", "DEP", "Críticos", "Taxa"],
        [[c["company"], c["movements"], c["arr"], c["dep"], c["critical"], f"{c['rate'] * 100:.1f}%"] for c in t["companies"]], 15)
    body += "<p class='note'>Relatório gerado automaticamente a partir dos dados carregados no FlightTracker. A timeline interactiva e o detalhe completo de pares estão disponíveis na aplicação e na exportação Excel.</p>"
    return wrap("Relatório de Análise de Dados", f"FlightTracker · Separação entre chegadas · {airport} · {time.strftime('%d/%m/%Y')}", body, landscape=True)


def report_daily(dy: dict, airport: str, blocks: list[str]) -> str:
    k = dy["kpis"]; n = dy["narrative"]; st = dy["stats"]
    body = "<h2>Filtros</h2>" + table(["Filtro", "Valor"], [["Período", ", ".join(s["dateFmt"] for s in st)], ["Limite crítico", f"{dy['limit']} s"],
                                                           ["Blocos horários", ", ".join(blocks) if blocks else "Todos"], ["Dias analisados", k["days"]]])
    body += "<h2>Síntese executiva</h2>" + "".join(f"<p>{esc(p)}</p>" for p in n["paragraphs"])
    body += "<h2>Destaques</h2><div class='insights'>" + "".join(f"<span>{esc(i)}</span>" for i in n["insights"]) + "</div>"
    body += "<h2>Indicadores principais</h2>" + kpi_grid([("Dias", k["days"]), ("Movimentos ARR+DEP", k["movs"]), ("Aterragens", k["arr"]), ("Descolagens", k["dep"]),
                                                          ("Pares", k["pairs"]), ("Abaixo do limite", k["crit"]), ("Críticos com DEP", k["critDep"]),
                                                          ("Pico aterragens", f"{k['peakArr']['hour']}h ({k['peakArr']['n']})" if k["peakArr"] else "—"),
                                                          ("Pico descolagens", f"{k['peakDep']['hour']}h ({k['peakDep']['n']})" if k["peakDep"] else "—"), ("Callsigns comuns", k["common"])])
    if st:
        top = max(st, key=lambda s: s["total"])
        body += f"<h2>Movimentos por dia</h2><p class='note'>Comparação de aterragens (ARR) e descolagens (DEP) por dia seleccionado. O dia {top['dateFmt']} destaca-se com {top['total']} movimentos (média global: {k['movs'] / max(1, k['days']):.1f}/dia).</p>"
        body += svg_grouped([("ARR", [s["arr"] for s in st], "#2563eb"), ("DEP", [s["dep"] for s in st], "#dc2626")], [s["label"] for s in st])
        body += "<h2>Ocorrências críticas por dia</h2>" + svg_bars([s["crit"] for s in st], [s["label"] for s in st], color="#dc2626")
        body += "<h2>Pico horário por dia</h2><p class='note'>Hora com maior volume de movimentos (ARR+DEP) em cada dia — indica o período de maior densidade operacional diária.</p>"
        body += svg_bars([s["peakN"] for s in st], [f"{s['label']} {s['peakHour'] or ''}h" for s in st])
    if len(dy["weekday"]) >= 1:
        wd = dy["weekday"]; c = max(wd, key=lambda e: e["avgCrit"])
        body += f"<h2 class='pb'>Padrões por dia da semana</h2><p class='note'>Agrupamento por dia da semana com médias por dia (comparação justa quando há diferente número de segundas, terças, etc.). {c['name']} regista a maior média de críticos ({c['avgCrit']}/dia).</p>"
        body += "<h3>Média de movimentos por dia da semana</h3>" + svg_bars([e["avgMov"] for e in wd], [e["name"][:3] for e in wd])
        body += "<h3>Média de pares críticos por dia da semana</h3>" + svg_bars([e["avgCrit"] for e in wd], [e["name"][:3] for e in wd], color="#dc2626")
        body += table(["Dia", "Nº dias", "Mov. totais", "Méd mov/dia", "Críticos", "Méd crít/dia", "Taxa %", "Datas"],
                      [[e["name"], e["days"], e["mov"], e["avgMov"], e["crit"], e["avgCrit"], e["rate"], ", ".join(e["dates"])] for e in wd])
    body += "<h2 class='pb'>Anexos tabulares</h2><h3>Resumo por dia</h3>" + table(["Data", "Dia sem.", "ARR", "DEP", "Movimentos", "Pares", "Críticos"],
        [[s["dateFmt"], s["wdName"], s["arr"], s["dep"], s["total"], s["pairs"], s["crit"]] for s in st], 30)
    body += "<h3>Companhias</h3>" + table(["Companhia", "Movimentos", "ARR", "DEP", "Críticos", "Taxa"], [[c["company"], c["movements"], c["arr"], c["dep"], c["critical"], f"{c['rate'] * 100:.1f}%"] for c in dy["companies"]], 10)
    body += "<h3>Callsigns comuns</h3>" + table(["Callsign", "Nº dias", "Dias", "Movimentos", "Companhia"], [[c["csn"], c["count"], ", ".join(fmt_date(d) for d in c["days"]), c["movements"], c["company"]] for c in dy["common"]], 20)
    body += "<h3>Resumo por dia e hora (amostra)</h3>" + table(["Data", "Dia", "Hora", "ARR", "DEP", "Movimentos", "Pares", "Críticos"],
        [[h["dateFmt"], h["wd"], f"{h['hour']}h", h["arr"], h["dep"], h["total"], h["pairs"], h["crit"]] for h in dy["hourTable"]], 40)
    return wrap("Relatório de Análise Diária", f"FlightTracker · {airport} · {time.strftime('%d/%m/%Y')}", body, landscape=True)


def report_rrms(rr: dict, shown: list[dict], cat_label: str, rwy_label: str) -> str:
    from .analysis.ops import rrms_stats
    body = f"<p class='note'>Filtros activos: Categoria = {esc(cat_label)} · Pista = {esc(rwy_label)} · Total passagens mostradas: {len(shown)}</p>"
    body += "<p class='note'>Pontos de referência: " + " · ".join(f"{v['label']} ({v['dms']})" for v in rr["points"].values()) + "</p>"
    body += ("<h3>Metodologia</h3><p>Para cada aeronave DEP, o algoritmo determina a pista pela direcção de movimento (latitude crescente→RWY02, decrescente→RWY20), "
             "encontra o par de amostras radar consecutivas com menor distância perpendicular ao ponto de referência (raio máx: 0.40 NM) e interpola linearmente a altitude e IAS no instante exacto de passagem.</p>")
    for rwy, twy in (("02", "TWY U6"), ("20", "TWY P")):
        rows = [h for h in shown if h["rwy"] == rwy]
        if not rows: continue
        body += f"<h2>RWY {rwy} ({len(rows)} passagens) · Ponto {twy}</h2>"
        st = rrms_stats(rows)
        paras = [f"Na RWY {rwy}, foram registadas {len(rows)} passagens no ponto de referência 2400m da THR deslocada (valores obtidos por interpolação linear entre amostras radar consecutivas)."]
        if st:
            paras.append(f"A altitude interpolada variou entre {st['altMin']} ft e {st['altMax']} ft, com média de {st['altAvg']} ft. " +
                         ("A maioria das aeronaves ainda se encontrava a baixa altitude no ponto de referência." if st["altAvg"] < 500 else
                          "A maioria das aeronaves já se encontrava a altitude positiva, confirmando que estavam no ar ao passar o ponto."))
            if st["iasMin"] is not None:
                paras.append(f"A IAS interpolada variou entre {st['iasMin']} kt e {st['iasMax']} kt, com média de {st['iasAvg']} kt.")
        from collections import Counter
        cats = Counter(h["cat"] for h in rows); n = len(rows)
        paras.append("Distribuição por categoria: " + " · ".join(f"{c} — {cats.get(c, 0)} ({cats.get(c, 0) / n * 100:.0f}%)" for c in ("M", "H", "?")) + ".")
        hc = Counter(h["hour"] for h in rows); ph = max(hc, key=lambda x: (hc[x], x))
        paras.append(f"Hora de pico: {ph}h com {hc[ph]} passagens.")
        air = sum(1 for h in rows if h["airborne"]); gnd = n - air; pa = round(air / n * 100)
        paras.append(f"Estado no ponto 2400m: {air} aeronaves no ar ({pa}%) · {gnd} no solo ({100 - pa}%)." +
                     (" Todas estavam airborne ao passar o ponto — condição favorável para RRMS." if gnd == 0 else " ⚠ Algumas aeronaves ainda estavam em rolagem ao passar o ponto — analisar casos individuais."))
        body += "".join(f"<p>{esc(p)}</p>" for p in paras)
        hours = sorted(hc)
        body += "<h3>No ar vs Solo no ponto 2400m</h3>" + svg_grouped([("No ar", [sum(1 for h in rows if h["hour"] == x and h["airborne"]) for x in hours], "#16a34a"),
                                                                       ("Solo", [sum(1 for h in rows if h["hour"] == x and not h["airborne"]) for x in hours], "#f59e0b")], [f"{x}h" for x in hours])
        body += "<h3>Estatísticas por categoria</h3>"
        cat_rows = []
        for lbl, sel in (("Total", rows), ("M (Medium)", [h for h in rows if h["cat"] == "M"]), ("H (Heavy)", [h for h in rows if h["cat"] == "H"])):
            s = rrms_stats(sel)
            if s: cat_rows.append([lbl, s["n"], s["altMin"], s["altAvg"], s["altMax"], fnum(s["iasMin"]), fnum(s["iasAvg"]), fnum(s["iasMax"])])
        body += table(["Categoria", "N", "Alt. mín", "Alt. méd", "Alt. máx", "IAS mín", "IAS méd", "IAS máx"], cat_rows)
        body += "<h3>Distribuição horária</h3>" + svg_bars([hc[x] for x in hours], [f"{x}h" for x in hours])
        body += "<h3>Detalhe completo</h3>" + table(["Callsign", "WTC", "Tipo A/C", "RWY", "Status", "Data", "Hora", "Alt. (ft)", "IAS (kt)", "Dist. (NM)"],
            [[h["csn"], h["cat"], h["acType"], h["rwy"], "No ar" if h["airborne"] else "Solo", h["date"], h["tod"], h["altFt"], fnum(h["ias"]), h["distNM"]] for h in rows])
    return wrap("RRMS — Relatório de Estatísticas", f"Reduced Runway Separation Minima · LPPT · {time.strftime('%d/%m/%Y %H:%M')}", body, landscape=True)


def report_ia(history: list[dict], provider_name: str) -> str:
    body = f"<p class='note'>Fornecedor: {esc(provider_name)}</p>"
    for m in history:
        if m["role"] == "user":
            body += f"<div style='border-left:4px solid #2563eb;background:#eff6ff;padding:8px 12px;margin:10px 0'><b>▶ Pergunta</b><br>{esc(m['text'])}</div>"
        else:
            txt = "<br>".join(esc(l).replace("- ", "• ", 1) if l.strip().startswith("- ") else esc(l) for l in m["text"].split("\n"))
            body += f"<div style='border-left:4px solid #16a34a;background:#f0fdf4;padding:8px 12px;margin:10px 0'><b>◈ {esc(provider_name)}</b><br>{txt}"
            g = m.get("grafico")
            if g and g.get("dados"):
                body += f"<h3>{esc(g.get('titulo', 'Gráfico'))}</h3>" + table(["Item", "Valor"], [[d.get("nome"), d.get("valor")] for d in g["dados"]])
            body += "</div>"
    return wrap("Análise IA — FlightTracker", f"Exportado em {time.strftime('%d/%m/%Y %H:%M')}", body)
